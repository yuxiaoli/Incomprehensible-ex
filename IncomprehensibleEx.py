import os
import re
import subprocess
import sublime
import sublime_plugin
import threading


SETTINGS_FILENAME = "incomprehensibleex.sublime-settings"
SETTINGS_CHANGE_KEY = "incomprehensible_ex"
ANSI_ESCAPE = re.compile(r"\x1B(?:[@-Z\\-_]|\[[0-?]*[ -/]*[@-~])")

DOCLING_SCRIPT = """
from pathlib import Path
import sys
from docling.document_converter import DocumentConverter

input_path = Path(sys.argv[1])
output_path = Path(sys.argv[2])
result = DocumentConverter().convert(input_path)
document = getattr(result, "document", None)
if document is None:
    raise SystemExit("Docling conversion failed: no document produced")
content = document.export_to_text()
if not content.endswith("\\n"):
    content += "\\n"
output_path.write_text(content, encoding="utf-8")
"""

MARKITDOWN_SCRIPT = """
from pathlib import Path
import sys
from markitdown import MarkItDown

input_path = sys.argv[1]
output_path = Path(sys.argv[2])
result = MarkItDown().convert(input_path)
content = getattr(result, "text_content", None)
if not content:
    raise SystemExit("MarkItDown conversion failed: no content produced")
if not content.endswith("\\n"):
    content += "\\n"
output_path.write_text(content, encoding="utf-8")
"""

ANYDOC_SCRIPT = """
from pathlib import Path
import sys
import anydoc

input_path = sys.argv[1]
output_path = Path(sys.argv[2])
try:
    content = anydoc.to_markdown(input_path)
except Exception as e:
    raise SystemExit(str(e))
if not content:
    raise SystemExit("anydoc conversion failed: no content produced")
if not content.endswith("\\n"):
    content += "\\n"
output_path.write_text(content, encoding="utf-8")
"""

PYMUPDF_SCRIPT = """
from pathlib import Path
import sys
import fitz
doc = fitz.open(sys.argv[1])
text = chr(12).join([page.get_text() for page in doc])
if not text.endswith("\\n"):
    text += "\\n"
Path(sys.argv[2]).write_text(text, encoding="utf-8")
"""

PYTHON_DOCX_SCRIPT = """
from pathlib import Path
import sys
import docx
doc = docx.Document(sys.argv[1])
text = "\\n".join([p.text for p in doc.paragraphs])
if not text.endswith("\\n"):
    text += "\\n"
Path(sys.argv[2]).write_text(text, encoding="utf-8")
"""

PYTHON_PPTX_SCRIPT = """
from pathlib import Path
import sys
import pptx
prs = pptx.Presentation(sys.argv[1])
text = []
for slide in prs.slides:
    for shape in slide.shapes:
        if hasattr(shape, "text"):
            text.append(shape.text)
text_str = "\\n".join(text)
if not text_str.endswith("\\n"):
    text_str += "\\n"
Path(sys.argv[2]).write_text(text_str, encoding="utf-8")
"""

OPENPYXL_SCRIPT = """
from pathlib import Path
import sys
import openpyxl
wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
text = []
for sheet in wb.worksheets:
    for row in sheet.iter_rows(values_only=True):
        text.append("\\t".join([str(c) if c is not None else "" for c in row]))
text_str = "\\n".join(text)
if not text_str.endswith("\\n"):
    text_str += "\\n"
Path(sys.argv[2]).write_text(text_str, encoding="utf-8")
"""


PDFLY_SCRIPT = """
from pathlib import Path
import sys
import subprocess
import os

input_path = sys.argv[1]
output_path = Path(sys.argv[2])

# Use PYTHONIOENCODING to force subprocess python to output utf-8
env = os.environ.copy()
env["PYTHONIOENCODING"] = "utf-8"

result = subprocess.run(
    ["pdfly", "extract-text", input_path], 
    capture_output=True,
    env=env
)
if result.returncode != 0:
    raise SystemExit(result.stderr.decode("utf-8", errors="replace"))

text = result.stdout.decode("utf-8", errors="replace")
if not text.endswith("\\n"):
    text += "\\n"
output_path.write_text(text, encoding="utf-8")
"""

def plugin_loaded():
    IncomprehensibleEx.ensure_settings_loaded()
    if IncomprehensibleEx.fileSettings is not None:
        IncomprehensibleEx.fileSettings.clear_on_change(SETTINGS_CHANGE_KEY)
        IncomprehensibleEx.fileSettings.add_on_change(
            SETTINGS_CHANGE_KEY, IncomprehensibleEx.ensure_settings_loaded
        )


def plugin_unloaded():
    if IncomprehensibleEx.fileSettings is not None:
        IncomprehensibleEx.fileSettings.clear_on_change(SETTINGS_CHANGE_KEY)


class IncomprehensibleEx(sublime_plugin.EventListener):

    print("** Incomprehensible Extensions Started **")

    DEFAULT_EXTENSIONS = [
        "mp3", "odt", "ogg", "pdf", "pptx", "ps", "psv", "rtf", "tff", "tif",
        "tiff", "tsv", "wav", "xls", "xlsx", "doc", "docx", "eml", "epub"
    ]
    EDITABLE_EXTENSIONS = [
        "asciidoc", "beamer", "commonmark", "context", "docbook", "docx",
        "dokuwiki", "dzslides", "fb2", "haddock", "html", "html5", "icml",
        "latex", "man", "markdown", "markdown_github", "markdown_mmd",
        "markdown_phpextra", "markdown_strict", "mediawiki", "native", "odt",
        "opendocument", "opml", "org", "plain", "revealjs", "rst", "rtf", "s5",
        "slideous", "slidy", "texinfo", "textile"
    ]
    SUPPORTED_ENGINES = ("docling", "markitdown", "pandoc", "anydoc")
    DEFAULT_ENGINE = "docling"

    extensions = list(DEFAULT_EXTENSIONS)
    editable_extensions = list(EDITABLE_EXTENSIONS)
    editMode = False
    engines = {"default": DEFAULT_ENGINE}
    fileSettings = None
    thread = None

    @classmethod
    def ensure_settings_loaded(cls):
        settings = sublime.load_settings(SETTINGS_FILENAME)
        changed = False

        if not settings.has("extensions"):
            settings.set("extensions", list(cls.DEFAULT_EXTENSIONS))
            changed = True
        if not settings.has("edit_mode"):
            settings.set("edit_mode", cls.editMode)
            changed = True
        if not settings.has("engines"):
            settings.set("engines", {"default": cls.DEFAULT_ENGINE})
            changed = True

        raw_extensions = settings.get("extensions", list(cls.DEFAULT_EXTENSIONS))
        if not isinstance(raw_extensions, list):
            raw_extensions = list(cls.DEFAULT_EXTENSIONS)
            settings.set("extensions", raw_extensions)
            changed = True

        normalized_extensions = []
        for item in raw_extensions:
            if isinstance(item, str) and item.strip():
                normalized_extensions.append(item.strip().lower())

        if not normalized_extensions:
            normalized_extensions = list(cls.DEFAULT_EXTENSIONS)
            settings.set("extensions", normalized_extensions)
            changed = True

        cls.extensions = normalized_extensions
        cls.editMode = bool(settings.get("edit_mode", False))

        raw_engines = settings.get("engines", {"default": cls.DEFAULT_ENGINE})
        cls.engines = {}
        
        if isinstance(raw_engines, dict):
            engine = raw_engines.get("default", cls.DEFAULT_ENGINE)
            if not isinstance(engine, str):
                engine = cls.DEFAULT_ENGINE
            
            for k, v in raw_engines.items():
                if k != "default" and isinstance(k, str) and isinstance(v, str):
                    cls.engines[k.lower()] = v.lower()
        else:
            engine = raw_engines if isinstance(raw_engines, str) else cls.DEFAULT_ENGINE

        engine = engine.strip().lower()
        if engine not in cls.SUPPORTED_ENGINES:
            print(
                "Incomprehensible Ex | Unsupported default engine '{0}', falling back to '{1}'".format(
                    engine, cls.DEFAULT_ENGINE
                )
            )
            engine = cls.DEFAULT_ENGINE
            
            if not isinstance(raw_engines, dict):
                settings.set("engines", {"default": engine})
                changed = True

        cls.engines["default"] = engine
        cls.fileSettings = settings

        if changed:
            sublime.save_settings(SETTINGS_FILENAME)

    @classmethod
    def get_proper_extension(cls, ext):
        engine = cls.engines.get(ext, cls.engines.get("default", cls.DEFAULT_ENGINE))
        if engine in ["docling", "markitdown", "anydoc", "pandoc"]:
            return ".md"
        return ".txt"

    def on_load(self, view):
        try:
            type(self).ensure_settings_loaded()
            file_info = self.get_file_info(view)
            if file_info is None:
                return

            path, file_name, extension = file_info
            if extension not in type(self).extensions:
                return

            window = view.window() or sublime.active_window()
            if window is not None:
                window.run_command("close")

            inp = os.path.join(path, file_name)
            self.thread = threading.Thread(
                target=self.handle_active,
                args=(view, inp, extension),
                name=file_name
            )
            self.thread.start()
        except Exception as error:
            print("Incomprehensible Ex on_load error:", error)

    def on_text_command(self, view, command_name, args):
        if command_name == "save" and view.settings().get("inex_original_file"):
            return ("inex_save", {})
        return None

    def on_modified(self, view):
        if view.settings().get("inex_original_file"):
            if view.is_scratch():
                view.set_scratch(False)

    def get_file_info(self, view):
        file_path = view.file_name()
        if not file_path:
            return None

        path, file_name = os.path.split(file_path)
        extension = os.path.splitext(file_name)[1].lstrip(".").lower()
        return path, file_name, extension

    @classmethod
    def get_python_launchers(cls):
        if os.name == "nt":
            return [
                ["py", "-3"],
                ["py"],
                ["python"],
                ["python3"],
                ["py", "-3.11"],
                ["py", "-3.12"],
            ]
        return [["python3"], ["python"]]

    @classmethod
    def build_command_specs(cls, inp, temp_out, ext, save):
        if save:
            return [
                {
                    "command": ["pandoc", "-s", "-o", temp_out, "-w", ext, inp],
                    "label": "pandoc",
                }
            ]

        specs = []
        
        specific_engine = cls.engines.get(ext)
        if specific_engine:
            specs.extend(cls.get_engine_specs(specific_engine, inp, temp_out))
            
        general_engine = cls.engines.get("default", cls.DEFAULT_ENGINE)
        if general_engine != specific_engine:
            specs.extend(cls.get_engine_specs(general_engine, inp, temp_out))
            
        return specs

    @classmethod
    def get_engine_specs(cls, engine, inp, temp_out):
        launchers = cls.get_python_launchers()
        if engine == "pandoc":
            return [
                {
                    "command": ["pandoc", "-s", "-t", "plain", "-o", temp_out, inp],
                    "label": "pandoc",
                }
            ]
        if engine == "docling":
            return [
                {
                    "command": launcher + ["-c", DOCLING_SCRIPT, inp, temp_out],
                    "label": "docling ({0})".format(" ".join(launcher)),
                }
                for launcher in launchers
            ]
        if engine == "markitdown":
            return [
                {
                    "command": launcher + ["-c", MARKITDOWN_SCRIPT, inp, temp_out],
                    "label": "markitdown ({0})".format(" ".join(launcher)),
                }
                for launcher in launchers
            ]
        if engine == "anydoc":
            return [
                {
                    "command": launcher + ["-c", ANYDOC_SCRIPT, inp, temp_out],
                    "label": "anydoc ({0})".format(" ".join(launcher)),
                }
                for launcher in launchers
            ]
        if engine == "pymupdf":
            return [
                {
                    "command": launcher + ["-c", PYMUPDF_SCRIPT, inp, temp_out],
                    "label": "pymupdf ({0})".format(" ".join(launcher)),
                }
                for launcher in launchers
            ]
        if engine == "pdfly":
            return [
                {
                    "command": launcher + ["-c", PDFLY_SCRIPT, inp, temp_out],
                    "label": "pdfly ({0})".format(" ".join(launcher)),
                }
                for launcher in launchers
            ]
        if engine == "python-docx":
            return [
                {
                    "command": launcher + ["-c", PYTHON_DOCX_SCRIPT, inp, temp_out],
                    "label": "python-docx ({0})".format(" ".join(launcher)),
                }
                for launcher in launchers
            ]
        if engine == "python-pptx":
            return [
                {
                    "command": launcher + ["-c", PYTHON_PPTX_SCRIPT, inp, temp_out],
                    "label": "python-pptx ({0})".format(" ".join(launcher)),
                }
                for launcher in launchers
            ]
        if engine == "openpyxl":
            return [
                {
                    "command": launcher + ["-c", OPENPYXL_SCRIPT, inp, temp_out],
                    "label": "openpyxl ({0})".format(" ".join(launcher)),
                }
                for launcher in launchers
            ]
            
        return [{"command": [engine, inp], "label": engine}]

    @classmethod
    def run_command_specs(cls, command_specs):
        startupinfo = None
        if os.name == "nt":
            startupinfo = subprocess.STARTUPINFO()
            startupinfo.dwFlags |= subprocess.STARTF_USESHOWWINDOW

        best_error_result = None
        retryable_error_markers = (
            "No module named",
            "ModuleNotFoundError",
            "ImportError",
            "No suitable Python runtime found",
            "Python was not found",
            "Requested Python version",
            "Unable to create process using",
        )

        for command_spec in command_specs:
            try:
                proc = subprocess.Popen(
                    command_spec["command"],
                    stdout=subprocess.PIPE,
                    stderr=subprocess.PIPE,
                    shell=False,
                    startupinfo=startupinfo
                )
            except FileNotFoundError as error:
                result = {
                    "returncode": 127,
                    "stdout": b"",
                    "stderr": str(error).encode("utf-8", errors="replace"),
                    "label": command_spec["label"],
                }
                if best_error_result is None:
                    best_error_result = result
                continue

            stdout, stderr = proc.communicate()
            result = {
                "returncode": proc.returncode,
                "stdout": stdout,
                "stderr": stderr,
                "label": command_spec["label"],
            }

            if proc.returncode == 0:
                return result

            stderr_text = stderr.decode("utf-8", errors="replace")
            is_retryable = any(marker in stderr_text for marker in retryable_error_markers)

            if not is_retryable:
                return result

            # For retryable errors, prefer module errors over runtime missing errors
            if best_error_result is None:
                best_error_result = result
            else:
                best_stderr = best_error_result["stderr"].decode("utf-8", errors="replace")
                if ("No module named" not in best_stderr and "ModuleNotFoundError" not in best_stderr) and \
                   ("No module named" in stderr_text or "ModuleNotFoundError" in stderr_text):
                    best_error_result = result

        return best_error_result

    @classmethod
    def fallback_to_bytes(cls, inp, temp_out):
        import shutil
        try:
            shutil.copyfile(inp, temp_out)
            return True
        except Exception as e:
            print("Failed to copy raw bytes:", e)
            return False

    @classmethod
    def convert(cls, inp, out, ext, save):
        try:
            temp_out = out + ".tmp"
            if os.path.exists(temp_out):
                os.remove(temp_out)
            if os.path.exists(out) and os.path.getsize(out) == 0:
                os.remove(out)

            command_specs = cls.build_command_specs(inp, temp_out, ext, save)
            result = cls.run_command_specs(command_specs)
            
            conversion_success = False
            
            if result is not None and result["returncode"] == 0:
                if not os.path.exists(temp_out) and result["stdout"]:
                    with open(temp_out, "wb") as f:
                        f.write(result["stdout"])
                if os.path.exists(temp_out):
                    conversion_success = True

            if not conversion_success:
                error_message = ""
                if result is not None:
                    error_message = (
                        result["stderr"].decode("utf-8", errors="replace").strip()
                        if result["stderr"] else "Unknown error."
                    )
                    error_message = ANSI_ESCAPE.sub("", error_message)
                    print("Incomprehensible Ex conversion failed")
                    print("Input:", inp)
                    print("Command:", result["label"])
                    print("Return code:", result["returncode"])
                    if error_message:
                        print(error_message)
                else:
                    print("No conversion command could be started.")

                print("Falling back to raw bytes display...")
                conversion_success = cls.fallback_to_bytes(inp, temp_out)
                
                if not conversion_success:
                    return False

            os.replace(temp_out, out)
            return True

        except Exception as error:
            print("Incomprehensible Ex exception:", error)
            if "temp_out" in locals():
                if cls.fallback_to_bytes(inp, temp_out):
                    try:
                        os.replace(temp_out, out)
                        return True
                    except Exception:
                        pass
            return False

    def handle_active(self, view, inp, ext):
        try:
            proper_ext = type(self).get_proper_extension(ext)
            temp_file = inp + proper_ext
            
            success = type(self).convert(inp, temp_file, ext, False)
            
            if success:
                try:
                    with open(temp_file, "r", encoding="utf-8") as f:
                        content = f.read()
                    os.remove(temp_file)
                except Exception as e:
                    print("Error reading converted file:", e)
                    return

                sublime.set_timeout(lambda: self.create_view(view.window(), inp, content, proper_ext), 0)
        except Exception as error:
            print("Incomprehensible Ex handle_active error:", error)

    def create_view(self, window, inp, content, proper_ext):
        if not window:
            window = sublime.active_window()
        if not window:
            return
            
        new_view = window.new_file()
        file_name = os.path.basename(inp)
        new_view.set_name(file_name + proper_ext)
        new_view.settings().set("inex_original_file", inp)
        new_view.settings().set("inex_proper_ext", proper_ext)
        
        if proper_ext == ".md":
            new_view.assign_syntax("Packages/Markdown/Markdown.sublime-syntax")
            
        new_view.run_command("append", {"characters": content})
        new_view.set_scratch(False)

class IncomprehensibleExEditModeOnCommand(sublime_plugin.ApplicationCommand):

    def run(self):
        try:
            IncomprehensibleEx.ensure_settings_loaded()
            IncomprehensibleEx.editMode = True
            IncomprehensibleEx.fileSettings.set('edit_mode', True)
            sublime.save_settings(SETTINGS_FILENAME)
            sublime.active_window().status_message("Incomprehensible Ex | Edit Mode ON")
        except Exception as e:
            print(e)

class IncomprehensibleExEditModeOffCommand(sublime_plugin.ApplicationCommand):

    def run(self):
        try:
            IncomprehensibleEx.ensure_settings_loaded()
            IncomprehensibleEx.editMode = False
            IncomprehensibleEx.fileSettings.set('edit_mode', False)
            sublime.save_settings(SETTINGS_FILENAME)
            sublime.active_window().status_message("Incomprehensible Ex | Edit Mode OFF")
        except Exception as e:
            print(e)


class InexSaveCommand(sublime_plugin.TextCommand):
    def run(self, edit):
        view = self.view
        inp = view.settings().get("inex_original_file")
        proper_ext = view.settings().get("inex_proper_ext", ".md")
        
        if not IncomprehensibleEx.editMode:
            sublime.status_message("Incomprehensible Ex: Edit Mode is OFF")
            return
            
        ext = os.path.splitext(inp)[1].lstrip(".").lower()
        if ext not in IncomprehensibleEx.editable_extensions:
            sublime.status_message(f"Incomprehensible Ex: {ext} is not editable")
            return

        content = view.substr(sublime.Region(0, view.size()))
        
        sublime.status_message(f"Incomprehensible Ex: Saving to {os.path.basename(inp)}...")
        
        temp_file = inp + proper_ext
        try:
            with open(temp_file, "w", encoding="utf-8") as f:
                f.write(content)
            
            success = IncomprehensibleEx.convert(temp_file, inp, ext, True)
            if success:
                sublime.status_message(f"Incomprehensible Ex: Saved {os.path.basename(inp)}")
                view.set_scratch(True)
            else:
                sublime.status_message(f"Incomprehensible Ex: Failed to save {os.path.basename(inp)}")
        finally:
            if os.path.exists(temp_file):
                try:
                    os.remove(temp_file)
                except:
                    pass
